from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".txt", ".md", ".json", ".csv", ".pdf", ".docx", ".xlsx"}


def extract_text(file_name: str, data: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"định dạng chưa hỗ trợ: {suffix or '(không có đuôi)'}")
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".json":
        parsed = json.loads(data.decode("utf-8", errors="replace"))
        return json.dumps(parsed, ensure_ascii=False, indent=2)
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        return "\n".join(" | ".join(str(cell) for cell in row) for row in rows)
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    if suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks: list[str] = []
        for ws in wb.worksheets:
            blocks.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    blocks.append(" | ".join(vals))
        return "\n".join(blocks)
    raise AssertionError("unreachable")
