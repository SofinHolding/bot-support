from __future__ import annotations

import argparse
import asyncio
import hashlib
import re
import unicodedata
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from .config import get_settings
from .db import Database, apply_migrations
from .knowledge.models import ExtractedKnowledgeUnit
from .knowledge.repository import KnowledgeRepository
from .knowledge.service import KnowledgeService
from .ragflow.client import RagFlowClient

EXPECTED_HEADERS = ("Tên", "Khách thường hỏi", "Nội dung")


def _cell(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _slug(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text)
    ascii_text = "".join(ch for ch in ascii_text if not unicodedata.combining(ch))
    ascii_text = ascii_text.lower().replace("đ", "d")
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_text).strip("-")
    return slug[:120] or "knowledge"


def _questions(raw: str) -> list[str]:
    return [x.strip() for x in re.split(r"\s*[|;\n]+\s*", raw) if x.strip()]


def workbook_units(path: Path) -> tuple[str, list[ExtractedKnowledgeUnit]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if not wb.worksheets:
        raise ValueError("workbook không có sheet")
    ws = wb.worksheets[0]
    header = tuple(_cell(ws.cell(1, i).value) for i in range(1, 4))
    if header != EXPECTED_HEADERS:
        raise ValueError(f"header không đúng: {header!r}; cần {EXPECTED_HEADERS!r}")

    rows: list[tuple[int, str, str, str]] = []
    for row_number in range(2, ws.max_row + 1):
        title = _cell(ws.cell(row_number, 1).value)
        questions_raw = _cell(ws.cell(row_number, 2).value)
        content = _cell(ws.cell(row_number, 3).value)
        if not title and not questions_raw and not content:
            continue
        if not title or not content:
            raise ValueError(f"dòng {row_number}: Tên và Nội dung là bắt buộc")
        rows.append((row_number, title, questions_raw, content))

    seen: dict[str, str] = {}
    units: list[ExtractedKnowledgeUnit] = []
    for row_number, title, questions_raw, content in rows:
        key = f"support.{_slug(title)}"
        old_title = seen.get(key)
        if old_title is not None and old_title != title:
            suffix = hashlib.sha256(title.encode("utf-8")).hexdigest()[:8]
            key = f"{key}-{suffix}"
        seen[key] = title

        questions = _questions(questions_raw)
        keywords = list(dict.fromkeys([title, *questions]))[:64]
        units.append(
            ExtractedKnowledgeUnit(
                knowledge_key=key,
                scope_key="default",
                subject=title,
                intent="support",
                canonical_title=title,
                canonical_summary=("Customer phrasings: " + " | ".join(questions)) if questions else title,
                keywords=keywords,
                language="mixed",
                content=content,
                identity_confidence=1.0,
                metadata={
                    "source_ref": f"Sheet {ws.title}, dòng {row_number}",
                    "customer_questions": questions,
                },
            )
        )
    return ws.title, units


async def clear_governance(db: Database) -> None:
    async with db.transaction() as conn:
        await conn.execute(
            """
            DELETE FROM knowledge_sync_outbox;
            DELETE FROM knowledge_ragflow_sync;
            DELETE FROM knowledge_conflicts;
            DELETE FROM knowledge_versions;
            DELETE FROM knowledge_units;
            DELETE FROM knowledge_sources;
            """
        )


async def run(path: Path, replace: bool) -> None:
    settings = get_settings()
    db = Database(settings.DATABASE_URL)
    await db.open()
    await apply_migrations(db)
    try:
        sheet, units = workbook_units(path)
        if replace:
            if settings.RAGFLOW_ENABLED:
                if not settings.RAGFLOW_API_KEY or not settings.RAGFLOW_DATASET_ID:
                    raise RuntimeError("--replace với RAGFLOW_ENABLED=true cần RAGFLOW_API_KEY và RAGFLOW_DATASET_ID")
                ragflow = RagFlowClient(
                    base_url=settings.RAGFLOW_BASE_URL,
                    api_key=settings.RAGFLOW_API_KEY,
                    timeout=settings.RAGFLOW_TIMEOUT_SECONDS,
                )
                # Fail before touching Postgres if the derived index cannot be cleared safely.
                await ragflow.delete_all_documents(dataset_id=settings.RAGFLOW_DATASET_ID)
            await clear_governance(db)

        repo = KnowledgeRepository(db)
        service = KnowledgeService(
            repo=repo,
            extractor=None,
            identity_min_confidence=settings.KNOWLEDGE_IDENTITY_MIN_CONFIDENCE,
            source_dir=settings.KNOWLEDGE_SOURCE_DIR,
        )
        report = await service.ingest_units(
            file_name=path.name,
            units=units,
            source_priority=100,
            source_type="canonical_xlsx",
            uploaded_by="migration",
            effective_from=datetime(2026, 9, 29, tzinfo=UTC),
            metadata={
                "sheet": sheet,
                "canonical": True,
                "duplicate_mirror": "content/",
                "source_path": str(path),
            },
        )
        active = sum(1 for item in report.units if item.status.value == "active")
        print(f"source={path.name} sheet={sheet} units={len(units)} active={active} counts={report.counts}")
    finally:
        await db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import clean InterLink XLSX into Knowledge Governance")
    parser.add_argument("path", nargs="?", default="raw-data/2026-09-29_V1-nguon-sach.xlsx")
    parser.add_argument("--replace", action="store_true", help="clear Governance tables before importing")
    args = parser.parse_args()
    asyncio.run(run(Path(args.path), args.replace))


if __name__ == "__main__":
    main()
